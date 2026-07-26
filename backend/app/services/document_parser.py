"""从上传的教案/课件中提取纯文本。"""
from __future__ import annotations

import io
import re
import zipfile
from typing import Any


def extract_from_bytes(filename: str, data: bytes) -> dict[str, Any]:
    name = (filename or "upload").lower()
    if name.endswith(".txt") or name.endswith(".md"):
        text = data.decode("utf-8", errors="ignore")
        return {"kind": "text", "filename": filename, "text": text.strip(), "chars": len(text)}
    if name.endswith(".docx"):
        return {"kind": "lesson", "filename": filename, **_docx_text(data)}
    if name.endswith(".pptx"):
        return {"kind": "ppt", "filename": filename, **_pptx_text(data)}
    # 兜底当文本
    text = data.decode("utf-8", errors="ignore")
    return {"kind": "text", "filename": filename, "text": text.strip()[:20000], "chars": len(text)}


def _docx_text(data: bytes) -> dict[str, Any]:
    try:
        from docx import Document  # type: ignore

        doc = Document(io.BytesIO(data))
        parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        text = "\n".join(parts)
        return {"text": text[:30000], "chars": len(text), "paragraphs": len(parts)}
    except Exception:
        # zip + document.xml 兜底
        text = _xml_zip_text(data, "word/document.xml")
        return {"text": text[:30000], "chars": len(text), "paragraphs": text.count("\n") + 1}


def _pptx_text(data: bytes) -> dict[str, Any]:
    slides: list[str] = []
    try:
        from pptx import Presentation  # type: ignore

        prs = Presentation(io.BytesIO(data))
        for i, slide in enumerate(prs.slides, 1):
            chunks: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    t = shape.text.strip()
                    if t:
                        chunks.append(t)
            if chunks:
                slides.append(f"【第{i}页】\n" + "\n".join(chunks))
        text = "\n\n".join(slides)
        return {"text": text[:40000], "chars": len(text), "slide_count": len(slides), "slides": slides[:40]}
    except Exception:
        text = _xml_zip_text(data, "ppt/slides/slide")
        return {"text": text[:40000], "chars": len(text), "slide_count": text.count("【第"), "slides": []}


def _xml_zip_text(data: bytes, prefix: str) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = sorted(n for n in zf.namelist() if n.startswith(prefix) and n.endswith(".xml"))
            if not names and prefix == "ppt/slides/slide":
                names = sorted(n for n in zf.namelist() if n.startswith("ppt/slides/") and n.endswith(".xml"))
            out: list[str] = []
            for i, n in enumerate(names, 1):
                raw = zf.read(n).decode("utf-8", errors="ignore")
                # strip tags
                plain = re.sub(r"<[^>]+>", " ", raw)
                plain = re.sub(r"\s+", " ", plain).strip()
                if plain:
                    out.append(f"【第{i}页】\n{plain}" if "slide" in n else plain)
            return "\n\n".join(out)
    except Exception:
        return ""
